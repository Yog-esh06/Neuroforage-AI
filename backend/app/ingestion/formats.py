"""Non-PDF parsers: images (OCR), spreadsheets, DOCX, plain text / markdown."""
from __future__ import annotations

import io
import re
from pathlib import Path

import pandas as pd

from app.core.exceptions import IngestionError
from app.ingestion import images as img_mod
from app.ingestion.layout import (ParsedDocument, ParsedElement, ParsedPage, assign_sections, finalize_page_text,
                                  looks_like_list_item, looks_like_numbered_heading)
from app.ingestion.ocr import OCRUnavailable
from app.ingestion.tables import normalize_table, table_to_text
from app.utils.text import normalize_text


def parse_image(data: bytes) -> ParsedDocument:
    try:
        img = img_mod.decode_image(data)
    except ValueError as e:
        raise IngestionError(str(e)) from e
    try:
        page = img_mod.ocr_image_to_page(img, 1)
    except OCRUnavailable as e:
        raise IngestionError(f"Image requires OCR but no OCR engine is available: {e}") from e
    doc = ParsedDocument(kind="image", pages=[page])
    h, w = img.shape[:2]
    doc.meta = {"width_px": w, "height_px": h}
    assign_sections(doc.pages)
    finalize_page_text(page)
    doc.title = next((e.text for e in page.elements if e.type == "heading"), None)
    return doc


def parse_text(data: bytes, *, markdown: bool = False) -> ParsedDocument:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    text = normalize_text(text)
    page = ParsedPage(1, method="text")
    order = 0
    for block in re.split(r"\n{2,}", text):
        b = block.strip()
        if not b:
            continue
        md = re.match(r"^(#{1,6})\s+(.*)$", b) if markdown else None
        if md:
            page.elements.append(ParsedElement("heading", md.group(2).strip(), 1, level=len(md.group(1)), order=order))
        elif "\n" not in b and (looks_like_numbered_heading(b) or (b.isupper() and 3 < len(b) < 90)):
            page.elements.append(ParsedElement("heading", b, 1, level=2, order=order))
        elif looks_like_list_item(b):
            page.elements.append(ParsedElement("list", b, 1, order=order))
        else:
            page.elements.append(ParsedElement("paragraph", re.sub(r"\s*\n\s*", " ", b), 1, order=order))
        order += 1
    doc = ParsedDocument(kind="text", pages=[page])
    assign_sections(doc.pages)
    finalize_page_text(page)
    doc.title = next((e.text for e in page.elements if e.type == "heading"), None)
    return doc


def parse_spreadsheet(data: bytes, filename: str) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".csv":
            sheets = {"Sheet1": pd.read_csv(io.BytesIO(data), header=None, dtype=str, keep_default_na=False)}
        else:
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=str, keep_default_na=False)
    except Exception as e:
        raise IngestionError(f"Spreadsheet could not be parsed: {e}") from e
    doc = ParsedDocument(kind="spreadsheet")
    for i, (name, df) in enumerate(sheets.items(), start=1):
        page = ParsedPage(i, method="table-only")
        page.elements.append(ParsedElement("heading", str(name), i, level=1, order=0))
        norm = normalize_table(df.values.tolist(), assume_header=True)
        if norm:
            page.elements.append(ParsedElement("table", table_to_text(norm), i, data={**norm, "sheet": str(name)},
                                               confidence=1.0, order=1))
        doc.pages.append(page)
    if not any(p.elements and len(p.elements) > 1 for p in doc.pages):
        raise IngestionError("Spreadsheet contains no data")
    doc.meta = {"sheets": list(sheets.keys())}
    doc.title = Path(filename).stem
    assign_sections(doc.pages)
    for p in doc.pages:
        finalize_page_text(p)
    return doc


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        from docx import Document as Docx

        d = Docx(io.BytesIO(data))
    except Exception as e:
        raise IngestionError(f"DOCX could not be parsed: {e}") from e
    page = ParsedPage(1, method="text")
    order = 0
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, d)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style)
            if m or style == "Title":
                page.elements.append(ParsedElement("heading", text, 1, level=int(m.group(1)) if m else 1, order=order))
            elif "List" in style or looks_like_list_item(text):
                page.elements.append(ParsedElement("list", text, 1, order=order))
            else:
                page.elements.append(ParsedElement("paragraph", text, 1, order=order))
            order += 1
        elif tag == "tbl":
            tbl = Table(child, d)
            norm = normalize_table([[c.text for c in r.cells] for r in tbl.rows])
            if norm:
                page.elements.append(ParsedElement("table", table_to_text(norm), 1, data=norm, confidence=1.0, order=order))
                order += 1
    if not page.elements:
        raise IngestionError("DOCX contains no text")
    doc = ParsedDocument(kind="docx", pages=[page])
    cp = d.core_properties
    doc.meta = {k: v for k, v in {"title": cp.title, "author": cp.author, "subject": cp.subject}.items() if v}
    assign_sections(doc.pages)
    finalize_page_text(page)
    doc.title = (cp.title or "").strip() or next((e.text for e in page.elements if e.type == "heading"), None)
    return doc
